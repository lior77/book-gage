#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Import an indicator from a CSV/TSV/XLSX you downloaded yourself.

This is the path that always works.  PORDATA, INE and idealista all offer a
Download button; the tables themselves are rendered by JavaScript, so scraping
them is brittle, but the downloaded file is stable.

    # PORDATA -> Municipios -> pick the indicator -> Download (XLSX or CSV)
    python3 scripts/import_indicator.py pordata_idade_mediana.xlsx \
        --key median_age --label "גיל חציוני" --unit "שנים" --year 2021 \
        --name-col "Municipios" --value-col "2021" \
        --source "PORDATA, Idade mediana da população residente" \
        --url "https://www.pordata.pt/municipios/idade+mediana+da+populacao-6"

    python3 scripts/build.py && python3 scripts/checks.py

Only the 18 municipalities of the Porto district (or the 243 parishes, with
--level freguesia) are kept; anything else in the file is ignored.  Rows whose
name cannot be resolved are listed rather than guessed at.
"""
import argparse
import csv
import os
import re
import sys

import indicator_store as store


def read_rows(path):
    ext = os.path.splitext(path)[1].lower()
    if ext in (".csv", ".tsv", ".txt"):
        for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
            try:
                text = open(path, encoding=enc).read()
                break
            except UnicodeDecodeError:
                continue
        else:
            sys.exit("cannot decode %s" % path)
        sample = text[:8000]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
            delim = dialect.delimiter
        except csv.Error:
            delim = "\t" if ext == ".tsv" else ";" if sample.count(";") > sample.count(",") else ","
        rows = list(csv.reader(text.splitlines(), delimiter=delim))
        return rows
    if ext in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook
        except ImportError:
            sys.exit("pip install openpyxl to read .xlsx, or export the sheet as CSV")
        wb = load_workbook(path, data_only=True, read_only=True)
        ws = wb[wb.sheetnames[0]]
        return [["" if c is None else c for c in row]
                for row in ws.iter_rows(values_only=True)]
    sys.exit("unsupported file type %s" % ext)


def find_header(rows, name_col, value_col):
    """PORDATA sheets carry a few title lines before the real header."""
    for i, row in enumerate(rows[:40]):
        cells = [str(c).strip() for c in row]
        low = [c.lower() for c in cells]
        if name_col.lower() in low and value_col.lower() in low:
            return i, low.index(name_col.lower()), low.index(value_col.lower())
    sys.exit("could not find a header row containing %r and %r.\n"
             "first rows were:\n%s" % (name_col, value_col,
                                       "\n".join(str(r[:6]) for r in rows[:12])))


def to_number(cell):
    s = str(cell).strip().replace(" ", "").replace(" ", "")
    if s in ("", "-", "//", "x", "n.d.", "nd", "ND"):
        return None
    s = s.replace("€", "").replace("%", "")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".")
    else:
        s = s.replace(",", ".")
    s = re.sub(r"[^0-9.\-]", "", s)
    try:
        return float(s)
    except ValueError:
        return None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("path")
    ap.add_argument("--key", required=True,
                    help="app indicator key: median_age, median_income, "
                         "foreign_pct, crimes_per_1000, price_eur_m2, ...")
    ap.add_argument("--label", required=True, help="Hebrew label shown in the UI")
    ap.add_argument("--unit", default="")
    ap.add_argument("--year", type=int, required=True, help="reference year")
    ap.add_argument("--name-col", required=True)
    ap.add_argument("--value-col", required=True)
    ap.add_argument("--source", required=True, help="cited in the UI, be exact")
    ap.add_argument("--url", default="")
    ap.add_argument("--level", choices=("municipio", "freguesia"), default="municipio")
    ap.add_argument("--scale", type=float, default=1.0)
    ap.add_argument("--warning", default="",
                    help="caveat shown next to the value, e.g. that a crime "
                         "figure is total registered offences and not violent crime")
    args = ap.parse_args()

    rows = read_rows(args.path)
    h, ni, vi = find_header(rows, args.name_col, args.value_col)
    values, unresolved = {}, []
    fidx = store.freguesia_index() if args.level == "freguesia" else {}

    for row in rows[h + 1:]:
        if len(row) <= max(ni, vi):
            continue
        name = str(row[ni]).strip()
        if not name:
            continue
        val = to_number(row[vi])
        if val is None:
            continue
        if args.level == "municipio":
            resolved = store.resolve_municipio(name)
        else:
            hit = fidx.get(store.norm(name))
            resolved = hit[0] if hit else None
        if resolved:
            values[resolved] = round(val * args.scale, 4)
        elif re.search(r"[A-Za-zÀ-ÿ]", name):
            unresolved.append(name)

    if unresolved:
        print("%d rows ignored (not in the Porto district, or name not matched)"
              % len(unresolved))
        print("  e.g. %s" % ", ".join(unresolved[:8]))
    meta = {"label_he": args.label, "unit": args.unit, "reference_year": args.year,
            "source_he": args.source, "url": args.url,
            "imported_from": os.path.basename(args.path)}
    if args.warning:
        meta["warning_he"] = args.warning
    if args.level == "municipio":
        store.save(args.key, meta, municipios=values)
    else:
        store.save(args.key, meta, freguesias=values)
    return 0


if __name__ == "__main__":
    sys.exit(main())
