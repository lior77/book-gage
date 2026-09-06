#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Read an INE Censos 2021 table and derive the age indicators from it.

Produces, for every municipality and every freguesia the file covers:

    ageing_index      65+ ÷ 0-14 × 100          (INE's índice de envelhecimento)
    pct_65plus        share of residents 65+
    pct_0_14          share of residents under 15
    pop_growth_pct    2011 → 2021, when the table carries both years
    median_age        only when the bands are fine enough to support it

Where to get the file
---------------------
    https://censos.ine.pt  →  Resultados definitivos  →  Quadros / Dados
Download at freguesia level.  Any of INE's usual shapes works: the coded
export (N_INDIVIDUOS_0_14 …), the published quadros with a title block above
the header, or a CSV.  Nothing here is hard-coded to one layout.

How to run
----------
    # 1. see what is actually in the file before importing anything
    python3 scripts/import_censos.py censos2021.xlsx --inspect

    # 2. import (add --sheet / --header-row / --col-* only if step 1 got it wrong)
    python3 scripts/import_censos.py censos2021.xlsx

    # 3. rebuild
    python3 scripts/build.py && python3 scripts/checks.py

    # prove the script itself works, without any file:
    python3 scripts/import_censos.py --selftest


CONFIDENCE (see scripts/ingest_overpass.py for the three levels)
---------------------------------------------------------------
Counts read straight from an INE table are "verified": INE is the statutory
source and the row is what it published.  Anything this script *computes* is
labelled by how safe the computation is:

  ageing index, shares, growth   arithmetic on published counts → "verified"
  median age from age bands      interpolation → "approx", and only attempted
                                 when the band holding the median is narrow
                                 enough for it to mean anything.  With INE's
                                 broad 0-14 / 15-24 / 25-64 / 65+ split it is
                                 not, so the script refuses rather than
                                 inventing a number.
"""
import argparse
import json
import os
import re
import sys
import unicodedata
from datetime import date

import indicator_store as store

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
RAW = os.path.join(ROOT, "data", "raw")

# The widest band we will interpolate a median inside. INE's coarse split puts
# the median somewhere in 25-64; spreading 40 years linearly says nothing, so
# anything wider than this is refused.
MAX_MEDIAN_BAND = 15


def norm(s):
    s = unicodedata.normalize("NFD", str(s or "").lower())
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9+]+", " ", s)).strip()


# --------------------------------------------------------------- reading ---
def read_rows(path, sheet=None):
    ext = os.path.splitext(path)[1].lower()
    if ext in (".csv", ".txt", ".tsv"):
        import csv
        for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
            try:
                text = open(path, encoding=enc).read()
                break
            except UnicodeDecodeError:
                continue
        else:
            sys.exit("cannot decode %s" % path)
        sample = text[:8000]
        delim = ";" if sample.count(";") > sample.count(",") else \
                "\t" if "\t" in sample.splitlines()[0] else ","
        return [list(r) for r in csv.reader(text.splitlines(), delimiter=delim)], "(csv)"
    if ext in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook
        except ImportError:
            sys.exit("pip install openpyxl (or save the sheet as CSV)")
        wb = load_workbook(path, data_only=True, read_only=True)
        name = sheet or wb.sheetnames[0]
        if name not in wb.sheetnames:
            sys.exit("sheet %r not found. sheets: %s" % (name, wb.sheetnames))
        ws = wb[name]
        rows = [["" if c is None else c for c in r] for r in ws.iter_rows(values_only=True)]
        return rows, name
    sys.exit("unsupported file type %s — export the sheet as .xlsx or .csv" % ext)


def find_header(rows, forced=None):
    """INE's published tables carry a title block above the real header."""
    if forced is not None:
        return forced
    best, best_score = 0, -1
    for i, row in enumerate(rows[:60]):
        cells = [norm(c) for c in row]
        filled = sum(1 for c in cells if c)
        if filled < 3:
            continue
        score = filled
        # a header row names a geography and at least one count
        if any(re.search(r"freguesia|concelho|municipio|dicofre|codigo|geocod|designacao", c)
               for c in cells):
            score += 20
        if any(re.search(r"populacao|residente|individuos|total|hm\b|\d+\s*(a|-)\s*\d+", c)
               for c in cells):
            score += 20
        if score > best_score:
            best, best_score = i, score
    return best


# norm() has already lowercased and turned _ and - into single spaces, so the
# separator between the two numbers may be a bare space ("hm 0 4"), the word
# "a" ("0 a 4 anos"), or "ate".
AGE_PATTERNS = [
    # "0 - 4", "0 a 4 anos", "HM_0_4", "N_INDIVIDUOS_0_14"
    (re.compile(r"(?:^|\s)(\d{1,3})\s+(?:a\s+|ate\s+)?(\d{1,3})(?:\s|$)"), "range"),
    # "65 ou mais", "75+", "85 e mais anos"
    (re.compile(r"(?:^|\s)(\d{1,3})\s*(?:\+|ou mais|e mais|ou superior)"), "open"),
]


def detect_columns(header):
    """Return {role: index} plus the list of age bands found."""
    cols, bands = {}, []
    for i, raw in enumerate(header):
        h = norm(raw)
        if not h:
            continue
        if "code" not in cols and re.search(r"\bdicofre\b|\bgeocod\b|\bcod(igo)?\b|\bcode\b", h):
            cols["code"] = i
        if "name" not in cols and re.search(r"freguesia|designacao|geodsg|\bnome\b|\blocal\b", h) \
                and not re.search(r"cod", h):
            cols["name"] = i
        if "mun" not in cols and re.search(r"concelho|municipio", h) and not re.search(r"cod", h):
            cols["mun"] = i
        if "pop2011" not in cols and "2011" in h and re.search(r"populacao|residente|individuos|total|hm", h):
            cols["pop2011"] = i
        if "pop2021" not in cols and re.search(r"populacao|residente|individuos", h) \
                and not re.search(r"\d{1,3}\s*(a|-|_)\s*\d{1,3}|ou mais|\+", h) and "2011" not in h:
            cols["pop2021"] = i
        if "median_age" not in cols and re.search(r"idade mediana|mediana da idade", h):
            cols["median_age"] = i
        if "mean_age" not in cols and re.search(r"idade media\b", h):
            cols["mean_age"] = i
        if "ageing" not in cols and re.search(r"indice de envelhecimento", h):
            cols["ageing"] = i

        for rx, kind in AGE_PATTERNS:
            m = rx.search(h)
            if m:
                if kind == "range":
                    lo, hi = int(m.group(1)), int(m.group(2))
                else:
                    lo, hi = int(m.group(1)), 100
                if lo <= hi <= 120:
                    bands.append({"lo": lo, "hi": hi, "col": i, "header": str(raw)})
                break
    # a band listed twice (male + female columns) would double-count
    seen, uniq = set(), []
    for b in sorted(bands, key=lambda b: (b["lo"], b["hi"])):
        key = (b["lo"], b["hi"])
        if key in seen:
            continue
        seen.add(key)
        uniq.append(b)
    return cols, uniq


def num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(" ", "").replace(" ", "")
    if s in ("-", "//", "x", "n.d.", "nd", ".."):
        return None
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".")
    else:
        s = s.replace(",", ".")
    s = re.sub(r"[^0-9.\-]", "", s)
    try:
        return float(s)
    except ValueError:
        return None


# ------------------------------------------------------------ derivation ---
def band_totals(row, bands):
    """Sum the age bands into the two groups the ageing index needs."""
    young = old = total = 0.0
    got = False
    for b in bands:
        v = num(row[b["col"]]) if b["col"] < len(row) else None
        if v is None:
            continue
        got = True
        total += v
        if b["hi"] <= 14:
            young += v
        if b["lo"] >= 65:
            old += v
    return (young, old, total) if got else (None, None, None)


def median_from_bands(row, bands):
    """Linear interpolation inside the band that holds the median.

    Returns (value, note) or (None, why not). The result is an estimate, never
    a published figure — the caller marks it "approx".
    """
    usable = []
    for b in bands:
        v = num(row[b["col"]]) if b["col"] < len(row) else None
        if v is not None:
            usable.append((b["lo"], b["hi"], v))
    usable.sort()
    total = sum(v for _, _, v in usable)
    if total <= 0:
        return None, "אין ספירות בפסי הגיל"
    half, cum = total / 2.0, 0.0
    for lo, hi, v in usable:
        if cum + v >= half:
            width = (hi - lo) + 1
            if width > MAX_MEDIAN_BAND:
                return None, ("החציון נופל בפס של %d שנים (%d–%d). אינטרפולציה על פס "
                              "כה רחב לא אומרת דבר, ולכן לא חושב ערך." % (width, lo, hi))
            if v <= 0:
                return None, "הפס שמכיל את החציון ריק"
            return round(lo + ((half - cum) / v) * width, 1), None
        cum += v
    return None, "לא נמצא פס שמכיל את החציון"


# -------------------------------------------------------------- matching ---
def resolve(row, cols, fidx):
    """Map an INE row onto a project key, by DICOFRE first and name second."""
    code = str(row[cols["code"]]).strip() if "code" in cols and cols["code"] < len(row) else ""
    code = re.sub(r"\D", "", code)
    name = str(row[cols["name"]]).strip() if "name" in cols and cols["name"] < len(row) else ""
    munname = str(row[cols["mun"]]).strip() if "mun" in cols and cols["mun"] < len(row) else ""

    if len(code) == 4:                                   # a municipality
        return ("municipio", store.INE_CODE.get(code)) if code in store.INE_CODE else (None, None)
    if len(code) == 6 and code[:4] in store.INE_CODE:    # a freguesia
        mun = store.INE_CODE[code[:4]]
        hit = fidx.get("%s %s" % (norm(mun), norm(name))) or fidx.get(norm(name))
        return ("freguesia", hit[0] if hit else None) if name else ("freguesia", None)
    if not code:                                         # name only
        mun = store.resolve_municipio(name)
        if mun and not munname:
            return "municipio", mun
        key = "%s %s" % (norm(munname), norm(name)) if munname else norm(name)
        hit = fidx.get(key) or fidx.get(norm(name))
        if hit:
            return "freguesia", hit[0]
    return None, None


# ------------------------------------------------------------------ main ---
def inspect(path, sheet, header_row):
    rows, sheetname = read_rows(path, sheet)
    h = find_header(rows, header_row)
    header = rows[h] if h < len(rows) else []
    cols, bands = detect_columns(header)
    print("sheet            %s" % sheetname)
    print("rows             %d" % len(rows))
    print("header row       %d  (0-based; override with --header-row)" % h)
    print("\ncolumns detected:")
    for role in ("code", "name", "mun", "pop2021", "pop2011", "median_age", "mean_age", "ageing"):
        if role in cols:
            print("  %-11s col %-3d  %s" % (role, cols[role], str(header[cols[role]])[:60]))
        else:
            print("  %-11s —" % role)
    print("\nage bands detected: %d" % len(bands))
    for b in bands:
        print("  %3d–%-3d  col %-3d  %s" % (b["lo"], b["hi"], b["col"], b["header"][:50]))
    if bands:
        wide = [b for b in bands if (b["hi"] - b["lo"] + 1) > MAX_MEDIAN_BAND]
        print("\nmedian age: estimated by interpolation and marked approx"
              + (", unless the median falls in one of the wide bands (%s), where "
                 "the script refuses rather than guess"
                 % ", ".join("%d–%d" % (b["lo"], b["hi"]) for b in wide) if wide else ""))
    print("\nfirst rows under the header:")
    for r in rows[h + 1:h + 4]:
        print("  " + " | ".join(str(c)[:18] for c in r[:8]))


def run(path, sheet, header_row, overrides, dry):
    rows, sheetname = read_rows(path, sheet)
    h = find_header(rows, header_row)
    header = rows[h]
    cols, bands = detect_columns(header)
    cols.update({k: v for k, v in overrides.items() if v is not None})
    if "name" not in cols and "code" not in cols:
        sys.exit("no geography column found — run with --inspect, then pass "
                 "--col-name / --col-code")

    fidx = store.freguesia_index()
    out = {k: {"municipios": {}, "freguesias": {}} for k in
           ("ageing_index", "pct_65plus", "pct_0_14", "pop_growth_pct", "median_age")}
    stats = {"rows": 0, "matched": 0, "unmatched": []}
    median_refusals = set()

    for row in rows[h + 1:]:
        if not any(str(c).strip() for c in row):
            continue
        stats["rows"] += 1
        level, key = resolve(row, cols, fidx)
        if not key:
            nm = str(row[cols.get("name", 0)]).strip() if cols.get("name") is not None else ""
            if nm:
                stats["unmatched"].append(nm)
            continue
        stats["matched"] += 1
        bucket = "municipios" if level == "municipio" else "freguesias"

        young, old, band_total = band_totals(row, bands)
        pop = num(row[cols["pop2021"]]) if "pop2021" in cols and cols["pop2021"] < len(row) else None
        pop = pop if pop else band_total

        if young is not None and old is not None:
            if young > 0:
                # arithmetic on published counts — as solid as the counts are
                out["ageing_index"][bucket][key] = round(old / young * 100, 1)
            if pop and pop > 0:
                out["pct_65plus"][bucket][key] = round(old / pop * 100, 1)
                out["pct_0_14"][bucket][key] = round(young / pop * 100, 1)

        if "pop2011" in cols and cols["pop2011"] < len(row):
            p11 = num(row[cols["pop2011"]])
            if p11 and pop:
                out["pop_growth_pct"][bucket][key] = round((pop - p11) / p11 * 100, 2)

        if "median_age" in cols and cols["median_age"] < len(row):
            v = num(row[cols["median_age"]])
            if v:
                out["median_age"][bucket][key] = round(v, 1)   # published, not estimated
        elif bands:
            v, why = median_from_bands(row, bands)
            if v:
                out["median_age"][bucket][key] = v
            elif why:
                median_refusals.add(why)

    print("sheet %s, header row %d" % (sheetname, h))
    print("rows read %d, matched to the district %d" % (stats["rows"], stats["matched"]))
    if stats["unmatched"]:
        print("ignored %d rows outside the Porto district (or unmatched), e.g. %s"
              % (len(stats["unmatched"]), ", ".join(stats["unmatched"][:5])))
    for why in median_refusals:
        print("median age: %s" % why)

    published_median = "median_age" in cols
    meta_common = {"reference_year": 2021, "url": "https://censos.ine.pt",
                   "imported_from": os.path.basename(path)}
    specs = {
        "ageing_index":   ("מדד הזדקנות", "65+ / 0-14 ×100", "verified",
                           "INE, Censos 2021 — מחושב מספירות שפורסמו: בני 65+ חלקי בני 0–14 כפול 100."),
        "pct_65plus":     ("אחוז בני 65+", "%", "verified",
                           "INE, Censos 2021 — מחושב מספירות שפורסמו."),
        "pct_0_14":       ("אחוז בני 0–14", "%", "verified",
                           "INE, Censos 2021 — מחושב מספירות שפורסמו."),
        "pop_growth_pct": ("שינוי אוכלוסייה 2011→2021", "%", "verified",
                           "INE — מחושב משני מפקדים שפורסמו."),
        "median_age":     ("גיל חציוני", "שנים",
                           "verified" if published_median else "approx",
                           "INE, Censos 2021 — ערך שפורסם." if published_median else
                           "אומדן: אינטרפולציה לינארית בתוך פס הגיל שמכיל את החציון. "
                           "לא ערך שפורסם, ולכן מסומן כמקורב."),
    }

    wrote = 0
    for k, (label, unit, conf, note) in specs.items():
        muns, freg = out[k]["municipios"], out[k]["freguesias"]
        if not muns and not freg:
            continue
        meta = dict(meta_common)
        meta.update({"label_he": label, "unit": unit, "source_he": note, "confidence": conf})
        if conf != "verified":
            meta["warning_he"] = note
        if dry:
            print("  would write %-16s municipalities %2d, freguesias %3d  [%s]"
                  % (k, len(muns), len(freg), conf))
        else:
            store.save(k, meta, municipios=muns or None, freguesias=freg or None)
        wrote += 1
    if not wrote:
        sys.exit("nothing could be derived — run --inspect and check the columns")
    return out if dry else 0


# -------------------------------------------------------------- selftest ---
def selftest():
    """Build a workbook shaped like an INE export and run the whole pipeline.

    Neither of us can see the real download yet, so this proves the detection,
    the arithmetic and the matching all work, and shows what the output looks
    like, without touching the project's data.
    """
    try:
        from openpyxl import Workbook
    except ImportError:
        sys.exit("pip install openpyxl to run the selftest")
    import tempfile

    wb = Workbook()
    ws = wb.active
    ws.title = "Q01"
    ws.append(["Censos 2021 — Resultados definitivos"])          # title block
    ws.append(["População residente por grupo etário"])
    ws.append([])
    bands = [(0, 4), (5, 9), (10, 14), (15, 19), (20, 24), (25, 29), (30, 34), (35, 39),
             (40, 44), (45, 49), (50, 54), (55, 59), (60, 64), (65, 69), (70, 74), (75, 100)]
    ws.append(["DICOFRE", "Concelho", "Freguesia", "População residente 2021",
               "População residente 2011"] + ["HM_%d_%d" % b for b in bands])

    # three real places, with invented counts: this exercises the code, it is
    # not data and never reaches data/raw
    fake = [
        ("130101", "Amarante", "Ansiães", 516, 601),
        ("131201", "Porto", "Bonfim", 22978, 24265),
        ("1312", "", "Porto", 231828, 237591),
    ]
    for code, mun, name, p21, p11 in fake:
        shape = [4, 5, 5, 5, 6, 6, 7, 7, 8, 8, 8, 8, 7, 6, 5, 5]
        s = sum(shape)
        counts = [round(p21 * x / s) for x in shape]
        counts[-1] += p21 - sum(counts)
        ws.append([code, mun, name, p21, p11] + counts)

    tmp = os.path.join(tempfile.mkdtemp(), "censos_fixture.xlsx")
    wb.save(tmp)
    print("selftest fixture: %s\n" % tmp)
    print("--- inspect " + "-" * 52)
    inspect(tmp, None, None)
    print("\n--- import (dry run, nothing is written) " + "-" * 25)
    out = run(tmp, None, None, {}, dry=True)

    print("\n--- checks " + "-" * 54)
    failures = []

    def check(label, got, want, tol=0.05):
        ok = got is not None and abs(got - want) <= tol
        print("  %-46s %-10s %s" % (label, got, "ok" if ok else "FAIL, expected %s" % want))
        if not ok:
            failures.append(label)

    # growth is independent of how the age bands were shaped
    check("Porto growth 2011→2021 (231828 vs 237591)",
          out["pop_growth_pct"]["municipios"].get("Porto"),
          round((231828 - 237591) / 237591 * 100, 2))
    check("Ansiães growth (516 vs 601)",
          out["pop_growth_pct"]["freguesias"].get("Amarante|Ansiães"),
          round((516 - 601) / 601 * 100, 2))
    # the fixture gives every row the same age shape, so the index is the same
    shape = [4, 5, 5, 5, 6, 6, 7, 7, 8, 8, 8, 8, 7, 6, 5, 5]
    want_idx = round(sum(shape[13:]) / sum(shape[:3]) * 100, 1)
    got = out["ageing_index"]["municipios"].get("Porto")
    print("  %-46s %-10s %s" % ("Porto ageing index (same shape for every row)",
          got, "ok" if got and abs(got - want_idx) < 2 else "FAIL, expected ~%s" % want_idx))
    if not (got and abs(got - want_idx) < 2):
        failures.append("ageing index")

    print("\nselftest %s. The counts are invented; the plumbing and the arithmetic "
          "are what was under test." % ("PASSED" if not failures else "FAILED: %s" % failures))
    return 1 if failures else 0


BANNER = (
    "\n  ┌─────────────────────────────────────────────────────────────┐\n"
    "  │  זהו סקריפט Python. מריצים אותו במחשב, בשורת הפקודה.        │\n"
    "  │  הוא *לא* שאילתת Overpass ואין להדביק אותו, או את הפלט      │\n"
    "  │  שלו, ב-overpass-turbo — שם זה ייתן parse error.            │\n"
    "  └─────────────────────────────────────────────────────────────┘\n")


def main():
    # Printed on every run: the Overpass queries and this script look alike from
    # the outside, and pasting either one into the other's tool fails obscurely.
    print(BANNER)
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("path", nargs="?", help="the INE .xlsx or .csv")
    ap.add_argument("--inspect", action="store_true", help="show the structure and stop")
    ap.add_argument("--selftest", action="store_true", help="run on a synthetic INE-shaped file")
    ap.add_argument("--sheet")
    ap.add_argument("--header-row", type=int)
    ap.add_argument("--dry-run", action="store_true")
    for role in ("code", "name", "mun", "pop2021", "pop2011", "median-age"):
        ap.add_argument("--col-" + role, type=int, help="0-based column index override")
    args = ap.parse_args()

    if args.selftest:
        return selftest()
    if not args.path:
        ap.error("give a file, or --selftest")
    if not os.path.exists(args.path):
        sys.exit("no such file: %s" % args.path)
    if args.inspect:
        inspect(args.path, args.sheet, args.header_row)
        return 0
    overrides = {"code": args.col_code, "name": args.col_name, "mun": args.col_mun,
                 "pop2021": args.col_pop2021, "pop2011": args.col_pop2011,
                 "median_age": args.col_median_age}
    # run() hands the derived values back on a dry run, for the selftest to
    # check; the process still exits 0.
    run(args.path, args.sheet, args.header_row, overrides, args.dry_run)
    return 0


if __name__ == "__main__":
    sys.exit(main())
